# -*- coding: utf-8 -*-
"""生成标准公文格式的通知 Word 文档（GB/T 9704-2012）"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

TITLE_FONT = "方正小标宋简体"
BODY_FONT = "仿宋"
OUT = r"C:\Users\cxx\WorkBuddy\Claw\通知作业_实训安全培训.docx"


def set_font(run, east, size, bold=False, ascii_font=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = ascii_font or east
    run._element.rPr.rFonts.set(qn('w:eastAsia'), east)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), ascii_font or east)


doc = Document()

# 页面设置 A4 + 公文页边距
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(3.7)
sec.bottom_margin = Cm(3.5)
sec.left_margin = Cm(2.8)
sec.right_margin = Cm(2.6)

style = doc.styles['Normal']
style.font.name = BODY_FONT
style.font.size = Pt(16)
style.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT)


def para(text="", font=BODY_FONT, size=16, bold=False, align=None,
         indent=True, right_indent=0, line=28.8):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line)
    if align is not None:
        p.alignment = align
    if indent:
        pf.first_line_indent = Pt(size * 2)
    if right_indent:
        pf.right_indent = Pt(right_indent)
    if text:
        r = p.add_run(text)
        set_font(r, font, size, bold)
    return p


C = WD_ALIGN_PARAGRAPH.CENTER
R = WD_ALIGN_PARAGRAPH.RIGHT
J = WD_ALIGN_PARAGRAPH.JUSTIFY

# 标题（二号小标宋，居中，两行梯形）
para("智能制造学院实训中心关于开展电气与工业互联网", TITLE_FONT, 22, False, C, indent=False, line=34)
para("实训安全专项整顿及规范操作培训的通知", TITLE_FONT, 22, False, C, indent=False, line=34)
para("", BODY_FONT, 16, False, None, indent=False, line=14)  # 标题下空一行

# 主送机关（顶格）
para("24级电气自动化、工业互联网技术专业全体学生：", BODY_FONT, 16, False, J, indent=False)

# 缘由段
para("近期，我院24级电气自动化、工业互联网技术专业学生在PLC实训、工业网关调试、产线模拟运维等实训环节中，存在违规带电操作、未佩戴绝缘防护用具、实训设备断电不彻底、实训工位杂物乱堆等安全隐患。为规范实训操作行为，杜绝安全事故，强化学生专业安全素养，经实训中心研究决定，开展全院电气与工业互联网实训安全专项整顿及规范操作培训。现将有关事项通知如下：", BODY_FONT, 16, False, J)

# 正文分条
blocks = [
    ("一、培训目的", [
        "进一步强化实训安全责任意识，规范电气与工业互联网实训操作流程，消除实训现场安全隐患，提高学生应急处置能力，切实保障实训教学安全有序开展。"]),
    ("二、培训时间", [
        "2026年9月15日（星期二）下午14:00—16:00。"]),
    ("三、培训地点", [
        "智能制造实训楼406工业互联网一体化实训车间。"]),
    ("四、参与人员", [
        "24级电气自动化、工业互联网技术专业全体在校生。"]),
    ("五、培训内容", [
        "（一）高低压电气实训安全规范；",
        "（二）工业物联网设备接线安全；",
        "（三）实训设备开关机流程；",
        "（四）应急触电处置方法；",
        "（五）实训工位6S管理要求。"]),
    ("六、相关要求", [
        "（一）全体参训学生须携带实训手册、签字笔，提前10分钟到达培训地点并签到；",
        "（二）培训期间须严格遵守现场纪律，认真听讲、做好记录，不得随意走动、交头接耳；",
        "（三）无特殊情况不得请假、缺席，不得迟到早退；确有特殊原因不能参加的，须提前向实训中心履行书面请假手续。"]),
]
for head, lines in blocks:
    para(head, BODY_FONT, 16, True, J)
    for ln in lines:
        para(ln, BODY_FONT, 16, False, J)

# 结语
para("特此通知。", BODY_FONT, 16, False, J)

# 落款（右空4字）
para("", BODY_FONT, 16, False, None, indent=False, line=14)
para("智能制造学院实训中心", BODY_FONT, 16, False, R, indent=False, right_indent=64)
para("2026年9月14日", BODY_FONT, 16, False, R, indent=False, right_indent=64)

doc.save(OUT)
print("saved:", OUT)
